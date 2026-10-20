"""Import World Bank commodity price indices (the "Pink Sheet") as world-level yearly tables.

Commodity prices hit countries through what they export and import: an oil-price crash hurts
oil exporters, a food-price spike hurts food importers. The forecast models combine these
world tables with each country's export and import shares (wdi-fuel-exports, wdi-food-imports...).

Tables use the entity WLD (the world), one row per year, in real 2010 US dollars (index 2010=100):
  pink-energy, pink-non-energy, pink-agriculture, pink-food, pink-grains, pink-fertilizers,
  pink-metals, pink-precious-metals, and pink-crude-oil (real $ per barrel, average of Brent,
  Dubai and WTI).

  mkdir -p .cache/pink-sheet && curl -L -o .cache/pink-sheet/CMO-Historical-Data-Annual.xlsx \\
      https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Annual.xlsx
  .cache/venv/bin/python tools/importers/pink_sheet.py      (needs openpyxl)

Licence: World Bank dataset terms, CC BY 4.0; see data/datasets/world-bank-commodity-prices.yaml.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import CACHE, write_series, write_table  # noqa: E402

FILE = CACHE / "pink-sheet" / "CMO-Historical-Data-Annual.xlsx"
INDICES = [  # column in "Annual Indices (Real)", table, title
    (2, "pink-energy", "Energy price index"), (3, "pink-non-energy", "Non-energy price index"),
    (4, "pink-agriculture", "Agriculture price index"), (6, "pink-food", "Food price index"),
    (8, "pink-grains", "Grains price index"), (13, "pink-fertilizers", "Fertilizer price index"),
    (14, "pink-metals", "Metals and minerals price index"), (16, "pink-precious-metals", "Precious metals price index"),
]
REGIONS = ["americas", "western-europe", "russia-eastern-europe", "middle-east-north-africa",
           "south-central-asia", "east-asia", "sub-saharan-africa", "south-east-asia"]


def rows_of(ws):
    return [r for r in ws.iter_rows(values_only=True) if isinstance(r[0], int)]


def main():
    import openpyxl

    wb = openpyxl.load_workbook(FILE, read_only=True, data_only=True)
    num = lambda v: float(v) if isinstance(v, (int, float)) else None
    out = [(col, table, title, "index (2010=100, real)", rows_of(wb["Annual Indices (Real)"])) for col, table, title in INDICES]
    prices = wb["Annual Prices (Real)"]
    header = next(r for r in prices.iter_rows(values_only=True) if r[1] == "Crude oil, average")
    assert header[1] == "Crude oil, average"
    out.append((1, "pink-crude-oil", "Crude oil price, average", "real 2010 US$ per barrel", rows_of(prices)))
    for col, table, title, unit, data in out:
        rows = [("WLD", str(r[0]), num(r[col])) for r in data]
        path = write_table(table, rows, digits=4)
        years = [int(r[1]) for r in rows if r[2] is not None]
        write_series({
            "id": table, "title": f"{title} (World Bank Pink Sheet)",
            "summary": f"World {title.lower()} in real 2010 US dollars, yearly, from the World Bank's commodity price data.",
            "occurred": {"start": str(min(years)), "end": str(max(years)), "precision": "year"},
            "threads": ["econ-crises"], "regions": REGIONS, "importance": 1, "unit": unit,
            "dataset": "world-bank-commodity-prices", "indicator": title, "table": f"tables/{table}.csv",
            "sources": [{"url": "https://www.worldbank.org/en/research/commodity-markets",
                         "title": "Commodity Markets (Pink Sheet data)", "publisher": "The World Bank", "type": "primary"}],
        })
        print(f"wrote {path.relative_to(CACHE.parent)}: {len(years)} years, {min(years)}-{max(years)}")


if __name__ == "__main__":
    main()
